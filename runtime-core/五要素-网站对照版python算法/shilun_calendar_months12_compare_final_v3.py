#!/usr/bin/env python3
"""V3 calendar calculation with AD 1 indexed month 3 as month-count zero.

The public month count is independently accumulated from AD 1 with remainder
46. Numerical five-element calculations retain the paired E1987 constants:
jiyue_e1987 = jiyue - 24566. This is an equivalent numbering change, not a
claim that the source literature prescribes an AD 1 astronomical epoch.
V2 is preserved. The default website comparison branch remains double_cal=0.

Reason for the change from V2: all 30000 cached month summaries have website
month count = V2 month count + 24566, while the shared daily arrays agree.
V3 aligns the public numbering without changing those numerical results.
Only month-state numbering, explicit E1987 conversion, input validation and
the distinct Excel output filename change; moon/solar corrections and
double_cal branch rules are retained. No application or database is migrated.
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


MOON_PERIODS = [7, 60, 60, 6, 707]
SUN_PERIODS = [27, 60, 60, 6, 67]
E1987_NUMBERED_MONTH_OFFSET = 12 * (1987 - 1)
E1987_ACCUMULATED_MONTH_OFFSET = 24566
AD1_INITIAL_REMAINDER = 46
DAY_LABELS = [f"{day}日" for day in range(1, 31)]


def reduction(values: Iterable[int], periods: Iterable[int]) -> list[int]:
    """Match the local MATLAB reduction function, including mixed-radix carry."""
    value = list(values)
    radix = list(periods)
    n = len(value)
    quotient = [0] * n
    remainder = [0] * n

    # The MATLAB function repeats this work n times. Its first pass reaches the
    # stable reduced state; preserving the loop makes its mutation semantics explicit.
    for _ in range(n):
        if value[-1] < 0:
            value[-1] += radix[-1]
            value[-2] -= 1
        remainder[-1] = value[-1] % radix[-1]
        quotient[-1] = math.floor(value[-1] / radix[-1])

        for index in range(n - 2, 0, -1):
            if value[index] < 0:
                value[index] += radix[index]
                value[index - 1] -= 1
            quotient[index] = math.floor((value[index] + quotient[index + 1]) / radix[index])
            remainder[index] = (value[index] + quotient[index + 1]) % radix[index]

        if value[0] < 0:
            value[0] += radix[0]
            remainder[0] = value[0]
        remainder[0] = (value[0] + quotient[1]) % radix[0]

    return remainder


def s_jiaoke(values: Iterable[int], periods: Iterable[int]) -> tuple[list[int], int]:
    """Match the MATLAB s_jiaoke quotient and final remainder calculation."""
    value = list(values)
    radix = list(periods)
    quotient = [0] * len(value)
    remainder = [0] * len(value)
    result = 0

    for _ in range(len(value)):
        remainder[0] = value[0] % 135
        quotient[0] = math.floor(value[0] / 135)
        for index in range(1, len(value)):
            quotient[index] = math.floor((value[index] + remainder[index - 1] * radix[index]) / 135)
            remainder[index] = (value[index] + remainder[index - 1] * radix[index]) % 135
            result = remainder[-1]
    return quotient, result


def calculate_month_state(year: int, month: int, double_cal: int) -> dict[str, object]:
    """Accumulate public month numbers from AD 1, indexed month 3."""
    if any(type(value) is not int for value in (year, month, double_cal)) or not (
        1 <= year <= 2500 and 1 <= month <= 12 and double_cal in (0, 1)
    ):
        raise ValueError("integer year 1..2500, month 1..12, double_cal 0/1 required")
    numbered_month = 12 * (year - 1) + (month + 12 if month < 3 else month) - 3
    carry, runyu_raw = divmod(2 * numbered_month + AD1_INITIAL_REMAINDER, 65)
    jiyue = numbered_month + carry
    runyu = runyu_raw
    mark_runyue = runyu_raw in (48, 49)
    if runyu_raw > 49 or (double_cal == 1 and mark_runyue):
        jiyue += 1
        runyu = runyu_raw + 2
    return {
        "jiyue0": numbered_month,
        "jiyue": jiyue,
        "jiyue_e1987": jiyue - E1987_ACCUMULATED_MONTH_OFFSET,
        "jiyue0_e1987": numbered_month - E1987_NUMBERED_MONTH_OFFSET,
        "runyu_raw": runyu_raw,
        "runyu": runyu,
        "mark_runyue": mark_runyue,
        "numbering_epoch": {"year": 1, "month": 3},
        "calculation_epoch": {"year": 1987, "month": 3},
    }


def calculate_month(year: int, month: int, double_cal: int) -> dict[str, object]:
    """Calculate the 30 daily six-element rows for one Tibetan calendar month."""
    state = calculate_month_state(year, month, double_cal)
    jiyue = int(state["jiyue_e1987"])
    yao_base = reduction([jiyue + 3, jiyue * 31 + 11, jiyue * 50 + 27, 2, jiyue * 480 + 332], MOON_PERIODS)
    sun_base = reduction([jiyue * 2, jiyue * 10, jiyue * 58, jiyue, jiyue * 17], SUN_PERIODS)
    rewo, qiaxi = reduction([jiyue * 2 + 21, jiyue + 90], [28, 126])

    moon_increment = [0, 59, 3, 4, 16]
    sun_increment = [0, 4, 21, 5, 43]
    moon_rough = [reduction(([0, 59, 3, 4, 16][index] + day * moon_increment[index] for index in range(5)), MOON_PERIODS) for day in range(30)]
    sun_rough = [reduction(([0, 4, 21, 5, 43][index] + day * sun_increment[index] for index in range(5)), SUN_PERIODS) for day in range(30)]
    mid_yao = [reduction([yao_base[index] + rough[index] for index in range(5)], MOON_PERIODS) for rough in moon_rough]
    mid_sun = [reduction([sun_base[index] + rough[index] for index in range(5)], SUN_PERIODS) for rough in sun_rough]
    t3 = [[math.floor(54 * day / 60), (54 * day) % 60] for day in range(1, 31)]

    rows = []
    for date in range(1, 31):
        quo_rewo = math.floor((rewo + date) / 14)
        rem_rewo = (rewo + date) % 14
        odd_level = quo_rewo % 2 != 0
        qian_jiao = rem_rewo <= 6
        moon_multipliers = [5, 5, 5, 4, 3, 2, 1, 1, 2, 3, 4, 5, 5, 5]
        moon_addends = [5, 10, 15, 19, 22, 24, 25, 24, 22, 19, 15, 10, 5, 0]
        moon_mul_num = moon_multipliers[rem_rewo]
        moon_add_num = moon_addends[13 if rem_rewo == 0 else rem_rewo - 1]

        moon_periods = [1, moon_mul_num, 60, 6, 707]
        moon_jinxinke = [moon_add_num]
        temp = qiaxi
        for radix in moon_periods[1:]:
            moon_jinxinke.append(math.floor(temp * radix / 126))
            temp = (temp * radix) % 126
        if qian_jiao:
            yao_jiaoke = reduction(
                [moon_jinxinke[0] + moon_jinxinke[1], *moon_jinxinke[2:]],
                [60, 60, 6, 707],
            )
        else:
            # Preserve the leading zero while borrowing across all time units.
            # Reducing the complemented tail separately loses one 707-unit carry
            # in the post-foot (hou-jiao) branch.
            yao_jiaoke = reduction(
                [
                    0,
                    moon_jinxinke[0] - moon_jinxinke[1],
                    -moon_jinxinke[2],
                    -moon_jinxinke[3],
                    -moon_jinxinke[4],
                ],
                [7, 60, 60, 6, 707],
            )[1:]
        yao_adjustment = [0, *yao_jiaoke]
        if odd_level:
            bandingyao = reduction([mid_yao[date - 1][index] - yao_adjustment[index] for index in range(5)], MOON_PERIODS)
        else:
            bandingyao = reduction([mid_yao[date - 1][index] + yao_adjustment[index] for index in range(5)], MOON_PERIODS)
        bandingyao_6 = [*bandingyao[:4], math.floor(bandingyao[4] * 67 / 707), (bandingyao[4] * 67) % 707]

        temp_rem0 = reduction([mid_sun[date - 1][0] - 6, mid_sun[date - 1][1] - 45, *mid_sun[date - 1][2:]], SUN_PERIODS)
        over_halfway = temp_rem0[0] > 13.5 or (temp_rem0[0] == 13 and temp_rem0[1] >= 30)
        temp_rem1 = reduction([temp_rem0[0] - 13, temp_rem0[1] - 30, *temp_rem0[2:]], SUN_PERIODS) if over_halfway else temp_rem0
        quo_s = math.floor((temp_rem1[1] + temp_rem1[0] * 60) / 135)
        rem_s = (temp_rem1[1] + temp_rem1[0] * 60) % 135
        if quo_s == 6:
            quo_s = 0
        sun_multipliers = [6, 4, 1, 1, 4, 6]
        sun_addends = [6, 10, 11, 10, 6, 0]
        sun_mul_num = sun_multipliers[quo_s]
        sun_add_num = sun_addends[5 if quo_s == 0 else quo_s - 1]
        quotient, _ = s_jiaoke([value * sun_mul_num for value in [rem_s, *temp_rem1[2:]]], [60, 60, 6, 67])
        sun_jinxinke = [sun_add_num, *quotient]
        if quo_s <= 2:
            sun_jiaoke0 = [sun_jinxinke[0] + sun_jinxinke[1], *sun_jinxinke[2:]]
        else:
            sun_jiaoke0 = [sun_jinxinke[0] - sun_jinxinke[1] - 1, *(value - adjustment for value, adjustment in zip([59, 5, 67], sun_jinxinke[2:]))]
        sun_jiaoke = reduction(sun_jiaoke0, [60, 60, 6, 67])

        if over_halfway:
            ding_yao = reduction([bandingyao_6[0], *(bandingyao_6[index + 1] + sun_jiaoke[index] for index in range(4)), bandingyao_6[5]], [7, 60, 60, 6, 67, 707])
            ding_sun = reduction([mid_sun[date - 1][0], *(mid_sun[date - 1][index + 1] + sun_jiaoke[index] for index in range(4))], SUN_PERIODS)
        else:
            ding_yao = reduction([bandingyao_6[0], *(bandingyao_6[index + 1] - sun_jiaoke[index] for index in range(4)), bandingyao_6[5]], [7, 60, 60, 6, 67, 707])
            ding_sun = reduction([mid_sun[date - 1][0], *(mid_sun[date - 1][index + 1] - sun_jiaoke[index] for index in range(4))], SUN_PERIODS)
        taiyin = [*reduction([ding_sun[0] + t3[date - 1][0], ding_sun[1] + t3[date - 1][1], *ding_sun[2:]], SUN_PERIODS), 0]
        taiyang = reduction([taiyin[0], *(taiyin[index + 1] - ding_yao[index + 1] for index in range(5))], [*SUN_PERIODS, 707])
        huihe = reduction([taiyang[index] + (ding_sun[index] if index < 5 else 0) for index in range(6)], [*SUN_PERIODS, 707])
        rows.append({"date": date, "ding_yao": ding_yao, "ding_sun": ding_sun, "taiyin": taiyin, "taiyang": taiyang, "huihe": huihe, "zuoyong": (2 * date - 1) % 7})

    return {**state, "yao_base": yao_base, "sun_base": sun_base, "rewo": rewo, "qiaxi": qiaxi, "moon_rough": moon_rough, "sun_rough": sun_rough, "t3": t3, "rows": rows}


def write_table(sheet, start_column: int, start_row: int, header: str, row_names: list[str], values: list[int]) -> None:
    sheet.cell(start_row, start_column + 1, header)
    for index, (row_name, value) in enumerate(zip(row_names, values), start=1):
        sheet.cell(start_row + index, start_column, row_name)
        sheet.cell(start_row + index, start_column + 1, value)


def write_column_table(sheet, start_column: int, start_row: int, header: str, values: list[int]) -> None:
    """Write a MATLAB table without RowNames in its starting column."""
    sheet.cell(start_row, start_column, header)
    for index, value in enumerate(values, start=1):
        sheet.cell(start_row + index, start_column, value)


def write_matrix_transposed(sheet, start_column: int, start_row: int, matrix: list[list[int]]) -> None:
    for row_index, row in enumerate(matrix):
        for column_index, value in enumerate(row):
            sheet.cell(start_row + column_index, start_column + row_index, value)


def export_year(year: int = 2027, double_cal: int = 0, output_dir: Path | str = Path(".")) -> Path:
    """Export all twelve months to the same workbook/sheet layout as MATLAB."""
    destination = Path(output_dir) / f"{year}年时轮历-check-v3-AD1.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for month in range(1, 13):
        result = calculate_month(year, month, double_cal)
        sheet = workbook.create_sheet(str(month))
        sheet.cell(5, 7, f"{year}年")
        sheet.cell(5, 8, f"{month}月")
        for column, label in enumerate(DAY_LABELS, start=17):
            sheet.cell(1, column, label)
        write_table(sheet, 6, 6, "积月闰余", ["积月", "闰余"], [int(result["jiyue"]), int(result["runyu"])])
        write_table(sheet, 7, 6, "曜基数", ["曜", "漏刻", "漏分", "呼吸", "707qiaxi"], result["yao_base"])
        write_table(sheet, 8, 6, "太阳基数", ["星宿 ", "漏刻", "漏分", "呼吸", "67qiaxi"], result["sun_base"])
        write_table(sheet, 9, 6, "整零数", ["整零数", ""], [int(result["rewo"]), int(result["qiaxi"])])
        write_matrix_transposed(sheet, 17, 2, result["moon_rough"])
        write_matrix_transposed(sheet, 17, 7, result["sun_rough"])
        write_matrix_transposed(sheet, 25, 15, result["t3"])

        for row in result["rows"]:
            start_row = 4 + 8 * row["date"]
            sheet.cell(start_row, 7, f"{row['date']}日六要素")
            write_table(sheet, 7, start_row + 1, "定曜", ["星宿", "漏刻", "漏分", "呼吸", "67qiaxi", "707qiaxi"], row["ding_yao"])
            write_table(sheet, 8, start_row + 1, "太阳日月宿", ["星宿", "漏刻", "漏分", "呼吸", "67qiaxi", "707qiaxi"], row["taiyang"])
            write_table(sheet, 9, start_row + 1, "定日", ["星宿 ", "漏刻", "漏分", "呼吸", "67qiaxi"], row["ding_sun"])
            write_column_table(sheet, 10, start_row + 1, "会合", row["huihe"])
            write_table(sheet, 11, start_row + 1, "太阴日月宿", ["星宿", "漏刻", "漏分", "呼吸", "67qiaxi", "707qiaxi"], row["taiyin"])

    workbook.save(destination)
    return destination


def open_workbook(path: Path | str):
    return load_workbook(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", type=int, default=2027)
    parser.add_argument("--double-cal", type=int, choices=(0, 1), default=0)
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    args = parser.parse_args()
    output = export_year(args.year, args.double_cal, args.output_dir)
    print(output)


if __name__ == "__main__":
    main()

